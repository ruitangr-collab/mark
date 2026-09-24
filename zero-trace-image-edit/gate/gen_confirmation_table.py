#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gen_confirmation_table.py — 参数化生成《局部改款·前置确认与验收确认表》Word
================================================================================
输入：confirmation(dict, 见 confirm.py 的 schema) + 可选 checks(闸门②结果列表)
输出：A4 Word，含 闸门①9项确认 / 闸门②9项点检 / 双签栏 / 水印。
水印：双签用户名+时间戳斜向淡印；未签时为占位文案（演示机制）。
依赖：python-docx（隔离 venv 已装 1.2.0）。
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from lxml import etree


# ---------- 字体/样式助手 ----------
def set_cjk(run, name="宋体", size=None, bold=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = etree.SubElement(rpr, qn("w:rFonts"))
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_para(doc, text="", size=10.5, bold=False, align=None, color=None, space_after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        set_cjk(p.add_run(text), size=size, bold=bold, color=color)
    return p


def set_cell(cell, text, size=9, bold=False, align=None, color=None):
    cell.text = ""
    for i, line in enumerate(str(text).split("\n")):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(1)
        if align is not None:
            p.alignment = align
        set_cjk(p.add_run(line), size=size, bold=bold, color=color)


def shade(cell, fill="4472C4"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="clear" w:color="auto" w:fill="{fill}"/>')
    tcPr.append(shd)


def hdr_row(table, headers, size=9):
    for i, h in enumerate(headers):
        set_cell(table.rows[0].cells[i], h, size=size, bold=True,
                 align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(table.rows[0].cells[i])


def style_table(table):
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False


def merge_row(table, row_idx, text, size=9, bold=False, fill="4472C4", color=RGBColor(0xFF, 0xFF, 0xFF)):
    cells = table.rows[row_idx].cells
    cells[0].merge(cells[-1])
    set_cell(cells[0], text, size=size, bold=bold, color=color)
    shade(cells[0], fill)


# ---------- 水印 ----------
def add_watermark(doc, text):
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0] if sec.header.paragraphs else sec.header.add_paragraph()
    wm = f'''<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
      xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">
      <v:shapetype id="_wm" coordsize="21600,21600" o:spt="136" adj="10800" path="m@4@5 l@4@11@9@11@9@5 xe" fillcolor="silver" stroked="f">
        <v:formulas><v:f eqn="prod @0 1 0"/><v:f eqn="sum @0 10402 0"/><v:f eqn="sum 20620 0 @1"/><v:f eqn="sum @2 0 @0"/><v:f eqn="sum @3 0 @0"/><v:f eqn="prod @0 2 1"/><v:f eqn="sum @6 0 @3"/><v:f eqn="sum @7 0 @3"/></v:formulas>
        <v:path textpathok="t" o:connecttype="rect"/><v:textpath on="t" fitshape="t"/></v:shapetype>
      <v:shape id="_wm1" type="#_wm" style="position:absolute;left:0;text-align:left;margin-left:0;margin-top:0;width:468pt;height:234pt;z-index:-251657216;mso-position-horizontal:center;mso-position-horizontal-relative:page;mso-position-vertical:center;mso-position-vertical-relative:page" rotation="315" fillcolor="silver" stroked="f">
        <v:textpath style="font-family:SimSun,Microsoft YaHei;font-size:40pt" o:extrusionok="f" fillcolor="silver" o:title="" string="{text}"/>
        <o:lock v:ext="edit" aspectratio="t"/></v:shape></w:pict>'''
    hp._p.append(parse_xml(wm))


# ---------- 主生成 ----------
def generate_confirmation_table(confirmation, checks=None, out_path="确认表.docx",
                                signed1=None, signed2=None):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Cm(29.7); sec.page_width = Cm(21.0)
    sec.top_margin = Cm(1.6); sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(1.8); sec.right_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"; normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    meta = confirmation.get("meta", {})
    it = confirmation.get("items", {})

    # 水印
    if signed1 and signed2 and signed1.get("username") and signed2.get("username"):
        wm = f"签名① {signed1['username']} {signed1.get('timestamp','')} · 签名② {signed2['username']} {signed2.get('timestamp','')}"
    else:
        wm = "PocketGraft 确认表 · 签名后自动生成双签水印"
    add_watermark(doc, wm)

    add_para(doc, "服装局部改款 · 前置确认与验收确认表", size=17, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x1F, 0x38, 0x64))
    add_para(doc, "PocketGraft 质量闸门　闸门①前置确认 ＋ 闸门②量化点检 ＋ 双签确认", size=10.5,
             align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x44, 0x72, 0xC4))
    add_para(doc, "模板版本 v2.0.0　|　未双签 = 禁止出图、禁止交付", size=9,
             align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x80, 0x80, 0x80), space_after=6)

    # 元信息
    m = doc.add_table(rows=3, cols=4); style_table(m)
    rows = [
        ("客户/品牌", meta.get("client", "____"), "款号/订单号", meta.get("style_no", "____")),
        ("目标款A", meta.get("image_a", "____"), "部件来源B", meta.get("image_b", "____")),
        ("操作人(登录用户名)", meta.get("operator", "____"), "出图日期", meta.get("date", "____")),
    ]
    for ri, (a, b, c, d) in enumerate(rows):
        set_cell(m.rows[ri].cells[0], a, size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF)); shade(m.rows[ri].cells[0])
        set_cell(m.rows[ri].cells[1], b)
        set_cell(m.rows[ri].cells[2], c, size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF)); shade(m.rows[ri].cells[2])
        set_cell(m.rows[ri].cells[3], d)
    add_para(doc, "", size=4, space_after=2)

    # 闸门①
    add_para(doc, "一、闸门①　前置确认（执行前，必须 9/9 逐项显式确认；不可默认、不可静默猜测）",
             size=11.5, bold=True, color=RGBColor(0x1F, 0x38, 0x64), space_after=3)
    g1 = doc.add_table(rows=1, cols=3); style_table(g1)
    hdr_row(g1, ["#", "确认主题", "使用者确认值（勾选/填写）"])

    def yn(v): return "☑" if v is True else ("☐" if v is False else "—")
    color = it.get("color", {})
    shape = it.get("shape", {})
    pos = it.get("position", {})
    stitch = it.get("stitching", {})
    g1_rows = [
        ("1", "运行环境（无外部依赖）", f"本地GPU: {yn(color and it.get('env',{}).get('local_gpu'))}　离线包: {yn(it.get('env',{}).get('offline'))}"),
        ("2", "颜色（料×色 3×3 矩阵）", f"料来源={color.get('fabric_src','?')}　色来源={color.get('color_src','?')}　勾选第{color.get('matrix_cell','?')}格"),
        ("3", "纹理方向", f"模式={it.get('texture',{}).get('mode','?')}（完全连续/近似对齐）"),
        ("4", "人话转译", f"方式={it.get('nlp',{}).get('mode','?')}（LLM解析/UI点选）"),
        ("5", "协调性容差", f"袋宽占比阈值≤{it.get('coord_tol',{}).get('pct','?')}%；生成前红蓝框预览"),
        ("6", "形状（子属性×来源）", f"袋型={shape.get('type','?')}　圆角R={shape.get('corner_r','?')}cm　袋盖={yn(shape.get('flap'))} 宽占{shape.get('flap_w_pct','?')}% 距口{shape.get('flap_dist_cm','?')}cm　铆钉={yn(shape.get('rivet'))} {shape.get('rivet_n','?')}颗　明线={shape.get('stitch','?')} {shape.get('stitch_mm','?')}mm　朝向锁袋口朝上"),
        ("7", "尺寸 SIZE", f"模式={it.get('size',{}).get('mode','?')}（B原/按比例/新）　{it.get('size',{}).get('w','?')}×{it.get('size',{}).get('h','?')}cm"),
        ("8", "位置（图上直摆·wearer_frame）", f"视图={pos.get('view','?')}　{pos.get('side','?')}（=着装者）　wearer=({pos.get('wearer_x','?')},{pos.get('wearer_y','?')})　回读人话={pos.get('human','?')}　镜像锁={yn(pos.get('mirror_locked'))}　镜像复制左片={yn(pos.get('mirror_copy_left'))}"),
        ("9", "缝纫线（子属性×来源）", f"线色={stitch.get('color_src','?')}　线型={stitch.get('type','?')}　线宽={stitch.get('width_mm','?')}mm　效果={stitch.get('effect','?')}"),
    ]
    for r in g1_rows:
        cells = g1.add_row().cells
        set_cell(cells[0], r[0], size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell(cells[1], r[1], size=9, bold=True)
        set_cell(cells[2], r[2], size=8.5)

    # 闸门②
    add_para(doc, "二、闸门②　交付前量化点检（逐项判定，任一 FAIL = 整体 FAIL）",
             size=11.5, bold=True, color=RGBColor(0x1F, 0x38, 0x64), space_after=3)
    g2 = doc.add_table(rows=1, cols=4); style_table(g2)
    hdr_row(g2, ["#", "检查项", "实测值", "判定"])
    if checks:
        for c in checks:
            cells = g2.add_row().cells
            set_cell(cells[0], str(c.get("idx", "")), size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell(cells[1], c.get("name", ""), size=9, bold=True)
            set_cell(cells[2], c.get("measured", ""), size=8.5)
            set_cell(cells[3], c.get("verdict", ""), size=8.5)
    else:
        g2.add_row()
        merge_row(g2, len(g2.rows) - 1, "（出图后由 checks.py 自动填入 9 项实测值与 PASS/FAIL）", size=9, fill="D9D9D9", color=RGBColor(0,0,0))
    add_para(doc, "总体核验：□ 9项全数PASS准予交付　□ 存在FAIL退回重做", size=10, bold=True, space_after=6)

    # 双签栏
    add_para(doc, "三、双签确认（采集：登录用户名自动带入 + 责任勾选 + 系统时间戳）",
             size=11.5, bold=True, color=RGBColor(0x1F, 0x38, 0x64), space_after=3)
    sg = doc.add_table(rows=6, cols=4); style_table(sg)
    merge_row(sg, 0, "签名① 意图确认（解锁出图）— 未签①禁止出图", size=10, fill="2E7D32")
    set_cell(sg.rows[1].cells[0], "操作人(登录用户名·自动带入)", size=9, bold=True)
    set_cell(sg.rows[1].cells[1], (signed1 or {}).get("username", "____"))
    set_cell(sg.rows[1].cells[2], "时间戳(系统)", size=9, bold=True)
    set_cell(sg.rows[1].cells[3], (signed1 or {}).get("timestamp", "____"))
    merge_row(sg, 2, f"□ 本人已逐项确认上述意图（未勾不许签）　责任勾选={yn((signed1 or {}).get('responsibility'))}", size=9, fill="FFFFFF", color=RGBColor(0,0,0))
    merge_row(sg, 3, f"是否已签：□ 已签（未签①=禁止出图）　哈希={ (signed1 or {}).get('hash','____')[:12] }", size=9, fill="FFFFFF", color=RGBColor(0,0,0))
    merge_row(sg, 4, "签名② 验收确认（解锁交付）— 未签②禁止交付", size=10, fill="2E7D32")
    set_cell(sg.rows[5].cells[0], "操作人(登录用户名·自动带入)", size=9, bold=True)
    set_cell(sg.rows[5].cells[1], (signed2 or {}).get("username", "____"))
    set_cell(sg.rows[5].cells[2], "时间戳(系统)", size=9, bold=True)
    set_cell(sg.rows[5].cells[3], (signed2 or {}).get("timestamp", "____"))

    note = doc.add_table(rows=1, cols=1); style_table(note)
    set_cell(note.rows[0].cells[0],
             "责任声明：签名①后冻结确认记录哈希，任一项改动即作废重走闸门①+签名①；签名②后发现问题须重走两道闸门。\n"
             "交付三件套：改款效果图（含3备选）+ 本双签水印确认表（Word）+ 点检报告，一并交付；缺件或无双签水印=禁止发出。",
             size=8.5, color=RGBColor(0x55, 0x55, 0x55))

    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    # 自测：用一份示例确认数据出一张表
    sample = {
        "meta": {"client": "示例客户A", "style_no": "DEMO-001",
                 "image_a": "A_denim.png", "image_b": "B_pocket.png", "operator": "示例操作员", "date": "2026-09-04"},
        "items": {
            "env": {"local_gpu": True, "offline": False},
            "color": {"fabric_src": "A", "color_src": "A", "matrix_cell": 1},
            "texture": {"mode": "approx"},
            "nlp": {"mode": "llm"},
            "coord_tol": {"pct": 40},
            "shape": {"type": "贴袋", "contour": "略梯形", "corner_r": 1.2, "flap": True,
                      "flap_w_pct": 60, "flap_dist_cm": 1.5, "rivet": False, "rivet_n": 0,
                      "rivet_gap_cm": 0, "stitch": "single", "stitch_mm": 3, "orientation_locked": True},
            "size": {"mode": "B", "w": 14, "h": 16},
            "position": {"view": "front", "side": "right", "wearer_x": 0.70, "wearer_y": 0.45,
                         "human": "右前片, 腰头下15cm, 侧缝内13cm, 袋口平行腰头",
                         "mirror_locked": True, "mirror_copy_left": False},
            "stitching": {"color_src": "accent", "type": "single", "width_mm": 3, "effect": "worn"},
        },
    }
    signed1 = {"username": "示例操作员", "timestamp": "2026-09-04 11:30", "responsibility": True, "hash": "abc123def456"}
    out = generate_confirmation_table(sample, checks=None, out_path="demo_确认表.docx",
                                       signed1=signed1, signed2=None)
    print("生成:", out, os.path.getsize(out), "bytes")

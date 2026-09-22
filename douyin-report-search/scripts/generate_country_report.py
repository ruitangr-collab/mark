#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
汇总20国抖音采集数据 → Excel 综合报告
"""
import json, sys, io
from pathlib import Path
from datetime import datetime
from collections import defaultdict

DATA_DIR = Path(r"C:/Users/53185/.workbuddy/skills/douyin-report-search/country_data")
OUTPUT   = DATA_DIR / "../抖音非洲20国采集汇总报告.xlsx"

def log(msg):
    sys.stdout.buffer.write(f"[{datetime.now().strftime('%H:%M:%S')}] {msg}\n".encode('utf-8'))
    sys.stdout.buffer.flush()

def load_all():
    """Load all country JSON files, return list of (country, keyword, data_dict)"""
    results = []
    for f in sorted(DATA_DIR.glob("*.json")):
        try:
            data = json.loads(f.read_text(encoding="utf-8"))
            country = data.get("country", "?")
            keyword = data.get("keyword", "?")
            results.append((country, keyword, data))
        except Exception as e:
            log(f"  SKIP {f.name}: {e}")
    return results

def main():
    log("Loading all country data...")
    all_data = load_all()
    log(f"  Loaded {len(all_data)} files")

    # ── Aggregate per country ──
    country_stats = {}
    keyword_records = []
    all_authors = {}

    for country, keyword, data in all_data:
        videos = data.get("videos", [])
        vcount = len(videos)
        
        # Unique authors in this keyword
        authors = {}
        total_likes = 0
        total_comments = 0
        total_shares = 0
        total_collects = 0
        
        for v in videos:
            auth = v.get("author", {})
            uid = auth.get("uid", "")
            nick = auth.get("nickname", "")
            if uid and uid not in authors:
                authors[uid] = {"nickname": nick, "uid": uid, "videos": 0, "total_likes": 0}
            if uid:
                authors[uid]["videos"] += 1
            
            stats = v.get("statistics", {})
            likes = stats.get("digg_count", 0) or 0
            comments = stats.get("comment_count", 0) or 0
            shares = stats.get("share_count", 0) or 0
            collects = stats.get("collect_count", 0) or 0
            total_likes += likes
            total_comments += comments
            total_shares += shares
            total_collects += collects
            
            if uid:
                authors[uid]["total_likes"] += likes
            
            # Track global authors
            if uid and uid not in all_authors:
                all_authors[uid] = {"nickname": nick, "uid": uid, "countries": set(), "keywords": set(), "total_videos": 0, "total_likes": 0}
            if uid:
                all_authors[uid]["countries"].add(country)
                all_authors[uid]["keywords"].add(keyword)
                all_authors[uid]["total_videos"] += 1
                all_authors[uid]["total_likes"] += likes

        avg_likes = int(total_likes / vcount) if vcount else 0
        avg_comments = int(total_comments / vcount) if vcount else 0
        
        keyword_records.append({
            "country": country,
            "keyword": keyword,
            "videos": vcount,
            "authors": len(authors),
            "total_likes": total_likes,
            "total_comments": total_comments,
            "total_shares": total_shares,
            "avg_likes": avg_likes,
            "avg_comments": avg_comments,
        })

        # Accumulate per country
        if country not in country_stats:
            country_stats[country] = {
                "total_videos": 0, "total_authors": {}, "total_likes": 0,
                "total_comments": 0, "total_shares": 0, "total_collects": 0,
                "keywords_done": 0, "keywords_total": 0,
            }
        cs = country_stats[country]
        cs["total_videos"] += vcount
        cs["total_likes"] += total_likes
        cs["total_comments"] += total_comments
        cs["total_shares"] += total_shares
        cs["total_collects"] += total_collects
        cs["keywords_total"] += 1
        if vcount > 0:
            cs["keywords_done"] += 1
        for uid, a in authors.items():
            if uid not in cs["total_authors"]:
                cs["total_authors"][uid] = a
            else:
                cs["total_authors"][uid]["videos"] += a["videos"]
                cs["total_authors"][uid]["total_likes"] += a["total_likes"]

    # ── Top authors globally ──
    top_authors_global = sorted(all_authors.values(), key=lambda x: x["total_videos"], reverse=True)[:50]
    top_authors_likes = sorted(all_authors.values(), key=lambda x: x["total_likes"], reverse=True)[:50]

    log(f"  Countries: {len(country_stats)}, Total videos: {sum(c['total_videos'] for c in country_stats.values())}")
    log(f"  Unique authors: {len(all_authors)}")

    # ── Generate Excel ──
    log("Generating Excel...")
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    wb = Workbook()
    header_font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align = Alignment(vertical="center")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )
    
    def style_header(ws, row, cols):
        for c in range(1, cols+1):
            cell = ws.cell(row=row, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

    def style_row(ws, row, cols):
        for c in range(1, cols+1):
            cell = ws.cell(row=row, column=c)
            cell.alignment = cell_align
            cell.border = thin_border
            cell.font = Font(name="微软雅黑", size=10)

    # ===== Sheet 1: 国家总览 =====
    ws1 = wb.active
    ws1.title = "国家总览"
    headers = ["排名", "国家", "视频总数", "关键词完成", "作者数", "总点赞", "总评论", "总分享", "单视频均赞", "单视频均评"]
    for c, h in enumerate(headers, 1):
        ws1.cell(row=1, column=c, value=h)
    style_header(ws1, 1, len(headers))

    sorted_countries = sorted(country_stats.items(), key=lambda x: x[1]["total_videos"], reverse=True)
    for rank, (country, cs) in enumerate(sorted_countries, 1):
        vc = cs["total_videos"]
        row_data = [
            rank, country, vc,
            f"{cs['keywords_done']}/{cs['keywords_total']}",
            len(cs["total_authors"]),
            cs["total_likes"], cs["total_comments"], cs["total_shares"],
            int(cs["total_likes"] / vc) if vc else 0,
            int(cs["total_comments"] / vc) if vc else 0,
        ]
        for c, val in enumerate(row_data, 1):
            ws1.cell(row=rank+1, column=c, value=val)
        style_row(ws1, rank+1, len(headers))

    ws1.column_dimensions["A"].width = 6
    ws1.column_dimensions["B"].width = 14
    ws1.column_dimensions["C"].width = 12
    ws1.column_dimensions["D"].width = 12
    ws1.column_dimensions["E"].width = 10
    ws1.column_dimensions["F"].width = 14
    ws1.column_dimensions["G"].width = 14
    ws1.column_dimensions["H"].width = 12
    ws1.column_dimensions["I"].width = 14
    ws1.column_dimensions["J"].width = 14

    # Number format for large numbers
    for row in range(2, len(sorted_countries)+2):
        for col in [6,7,8,9,10]:
            ws1.cell(row=row, column=col).number_format = "#,##0"

    # ===== Sheet 2: 关键词明细 =====
    ws2 = wb.create_sheet("关键词明细")
    headers2 = ["国家", "关键词", "视频数", "作者数", "总点赞", "总评论", "总分享", "均赞", "均评"]
    for c, h in enumerate(headers2, 1):
        ws2.cell(row=1, column=c, value=h)
    style_header(ws2, 1, len(headers2))

    keyword_records.sort(key=lambda x: (-x["videos"], x["country"], x["keyword"]))
    for i, kr in enumerate(keyword_records, 2):
        row_data = [kr["country"], kr["keyword"], kr["videos"], kr["authors"],
                     kr["total_likes"], kr["total_comments"], kr["total_shares"],
                     kr["avg_likes"], kr["avg_comments"]]
        for c, val in enumerate(row_data, 1):
            ws2.cell(row=i, column=c, value=val)
        style_row(ws2, i, len(headers2))

    for col in ["A","B"]: ws2.column_dimensions[col].width = 16
    for col in ["C","D"]: ws2.column_dimensions[col].width = 10
    for col in ["E","F","G","H","I"]: ws2.column_dimensions[col].width = 13

    for row in range(2, len(keyword_records)+2):
        for col in [3,4,5,6,7,8,9]:
            ws2.cell(row=row, column=col).number_format = "#,##0"

    # ===== Sheet 3: 全球热门作者TOP50(按视频数) =====
    ws3 = wb.create_sheet("热门作者TOP50")
    headers3 = ["排名", "作者昵称", "视频数", "总点赞", "覆盖国家数", "覆盖关键词"]
    for c, h in enumerate(headers3, 1):
        ws3.cell(row=1, column=c, value=h)
    style_header(ws3, 1, len(headers3))

    for i, a in enumerate(top_authors_global, 2):
        row_data = [i-1, a["nickname"], a["total_videos"], a["total_likes"],
                     len(a["countries"]), ", ".join(sorted(a["countries"]))]
        for c, val in enumerate(row_data, 1):
            ws3.cell(row=i, column=c, value=val)
        style_row(ws3, i, len(headers3))

    ws3.column_dimensions["A"].width = 6
    ws3.column_dimensions["B"].width = 22
    ws3.column_dimensions["C"].width = 10
    ws3.column_dimensions["D"].width = 14
    ws3.column_dimensions["E"].width = 12
    ws3.column_dimensions["F"].width = 40

    for row in range(2, len(top_authors_global)+2):
        for col in [3,4]:
            ws3.cell(row=row, column=col).number_format = "#,##0"

    # ===== Sheet 4: 各国头部作者 =====
    ws4 = wb.create_sheet("各国头部作者")
    headers4 = ["国家", "排名", "作者昵称", "视频数", "总点赞"]
    for c, h in enumerate(headers4, 1):
        ws4.cell(row=1, column=c, value=h)
    style_header(ws4, 1, len(headers4))

    row4 = 2
    for country, cs in sorted_countries:
        top_auth = sorted(cs["total_authors"].values(), key=lambda x: x["videos"], reverse=True)[:10]
        for rank, a in enumerate(top_auth, 1):
            row_data = [country, rank, a["nickname"], a["videos"], a["total_likes"]]
            for c, val in enumerate(row_data, 1):
                ws4.cell(row=row4, column=c, value=val)
            style_row(ws4, row4, len(headers4))
            row4 += 1

    ws4.column_dimensions["A"].width = 14
    ws4.column_dimensions["B"].width = 6
    ws4.column_dimensions["C"].width = 22
    ws4.column_dimensions["D"].width = 10
    ws4.column_dimensions["E"].width = 14

    for row in range(2, row4):
        for col in [4,5]:
            ws4.cell(row=row, column=col).number_format = "#,##0"

    # ===== Sheet 5: 空关键词 =====
    ws5 = wb.create_sheet("无结果关键词")
    headers5 = ["国家", "关键词", "状态"]
    for c, h in enumerate(headers5, 1):
        ws5.cell(row=1, column=c, value=h)
    style_header(ws5, 1, len(headers5))

    empty_kws = [kr for kr in keyword_records if kr["videos"] == 0]
    for i, kr in enumerate(empty_kws, 2):
        row_data = [kr["country"], kr["keyword"], "无数据"]
        for c, val in enumerate(row_data, 1):
            ws5.cell(row=i, column=c, value=val)
        style_row(ws5, i, len(headers5))

    ws5.column_dimensions["A"].width = 16
    ws5.column_dimensions["B"].width = 16
    ws5.column_dimensions["C"].width = 10

    # Save
    wb.save(OUTPUT)
    log(f"\nReport saved: {OUTPUT}")
    log(f"Sheets: 国家总览, 关键词明细, 热门作者TOP50, 各国头部作者, 无结果关键词")
    
    # Print summary
    log(f"\n{'='*60}")
    log("SUMMARY")
    log(f"{'='*60}")
    total_v = sum(c["total_videos"] for c in country_stats.values())
    log(f"Total videos: {total_v}")
    log(f"Total authors: {len(all_authors)}")
    log(f"Countries: {len(country_stats)}")
    log(f"Keywords with data: {sum(1 for kr in keyword_records if kr['videos'] > 0)}/{len(keyword_records)}")
    log(f"Keywords with 0 results: {len(empty_kws)}")
    for kr in empty_kws:
        log(f"  - {kr['country']}: {kr['keyword']}")

if __name__ == "__main__":
    main()

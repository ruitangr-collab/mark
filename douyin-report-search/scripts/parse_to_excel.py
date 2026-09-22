#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
抖音视频数据解析 → Excel 报告
从 _raw 字段直接提取：作者信息、粉丝数、互动数据、话题、评论
"""

import json, sys, re
from pathlib import Path
from datetime import datetime
from collections import Counter

WORK_DIR = Path(r"C:/Users/53185/.workbuddy/skills/douyin-report-search")
RAW_FILE = WORK_DIR / "douyin_raw_data.json"
PROFILE_CACHE = WORK_DIR / "author_profiles_cache.json"

def log(msg):
    sys.stdout.buffer.write(f"[{datetime.now().strftime('%H:%M:%S')}] {msg}\n".encode('utf-8'))
    sys.stdout.buffer.flush()

def num_str_to_int(s):
    if not s: return 0
    s = str(s).strip()
    if '万' in s:
        return int(float(s.replace('万','')) * 10000)
    if '亿' in s:
        return int(float(s.replace('亿','')) * 100000000)
    try: return int(re.sub(r'[^\d]','', s) or '0')
    except: return 0

def parse_video(v):
    """从采集数据解析完整视频信息"""
    raw = v.get("_raw", {})
    author = raw.get("author", {})
    stats = raw.get("statistics", {})
    cha_list = raw.get("cha_list", []) or []
    comment_list = raw.get("comment_list", []) or []

    # 话题标签
    tags = [c.get("cha_name","") for c in cha_list if c.get("cha_name")]

    # 热门评论（取前5条）
    top_comments = []
    for c in comment_list[:5]:
        u = c.get("user", {})
        text = c.get("text","")
        # 清理 emoji 避免 Excel 乱码
        text_clean = re.sub(r'[^\u0000-\uFFFF]', '', text)
        top_comments.append({
            "text": text_clean[:200],
            "digg_count": c.get("digg_count", 0),
            "nickname": u.get("nickname", ""),
        })

    # 从 desc 解析时长/播放量
    desc = v.get("desc", "")
    dur_m = re.match(r'^(\d+:\d+(?::\d+)?)\s*', desc)
    duration = dur_m.group(1) if dur_m else ""

    play_raw = ""
    play_m = re.search(r'^\d+:\d+(?::\d+)?\s+([\d.]+[万亿]?)\s+', desc)
    if play_m:
        play_raw = play_m.group(1)

    return {
        "视频ID": raw.get("aweme_id", v.get("aweme_id","")),
        "视频链接": f"https://www.douyin.com/video/{raw.get('aweme_id', v.get('aweme_id',''))}",
        "作者昵称": author.get("nickname", ""),
        "作者ID": author.get("sec_uid", ""),
        "粉丝数": author.get("follower_count", 0),
        "粉丝数(文本)": format_count(author.get("follower_count", 0)),
        "点赞数": stats.get("digg_count", 0),
        "评论数": stats.get("comment_count", 0),
        "分享数": stats.get("share_count", 0),
        "收藏数": stats.get("collect_count", 0),
        "下载数": stats.get("download_count", 0),
        "话题标签": " ".join(f"#{t}" for t in tags[:8]),
        "视频描述": clean_text(raw.get("desc", "")[:300]),
        "时长": duration,
        "播放量(文本)": play_raw,
        "来源关键词": v.get("_source_keyword", ""),
        "发布时间戳": raw.get("create_time", 0),
        "热评1": top_comments[0]["text"] if len(top_comments) > 0 else "",
        "热评1点赞": top_comments[0]["digg_count"] if len(top_comments) > 0 else 0,
        "热评2": top_comments[1]["text"] if len(top_comments) > 1 else "",
        "热评2点赞": top_comments[1]["digg_count"] if len(top_comments) > 1 else 0,
        "热评3": top_comments[2]["text"] if len(top_comments) > 2 else "",
        "热评3点赞": top_comments[2]["digg_count"] if len(top_comments) > 2 else 0,
    "作者主页": f"https://www.douyin.com/user/{author.get('sec_uid','')}",
    "作者签名": "",  # 稍后从缓存填充
    "_top_comments": top_comments,
}

def format_count(n):
    if n >= 100000000: return f"{n/100000000:.1f}亿"
    if n >= 10000: return f"{n/10000:.1f}万"
    return str(n)

def clean_text(s):
    """清理文本，去掉 emoji"""
    s = re.sub(r'[^\u0000-\uFFFF]', '', s)
    return s.strip()

def main():
    log("加载采集数据...")
    d = json.loads(RAW_FILE.read_text(encoding="utf-8"))
    videos_raw = d.get("videos", [])
    log(f"共 {len(videos_raw)} 条视频")

    # 加载作者签名缓存
    profiles = {}
    if PROFILE_CACHE.exists():
        profiles = json.loads(PROFILE_CACHE.read_text(encoding="utf-8"))
        sig_count = sum(1 for v in profiles.values() if v.get("signature"))
        log(f"加载作者资料: {len(profiles)} 人, {sig_count} 人有签名")

    rows = []
    for i, v in enumerate(videos_raw):
        row = parse_video(v)
        # 从缓存填入作者签名
        uid = row["作者ID"]
        if uid in profiles:
            sig = profiles[uid].get("signature", "") or ""
            row["作者签名"] = clean_text(sig)
        rows.append(row)
        if (i+1) % 50 == 0:
            log(f"  解析进度: {i+1}/{len(videos_raw)}")

    log(f"解析完成，共 {len(rows)} 条")

    # ── 生成作者汇总 ──
    # 用 ASCII key 存储，避免 Windows 编码问题
    author_stats = {}
    for row in rows:
        aid = row["作者ID"]
        if not aid: continue
        if aid not in author_stats:
            author_stats[aid] = {
            "nickname": row["作者昵称"],
            "sec_uid": aid,
            "signature": row["作者签名"],
            "follower_count": row["粉丝数"],
                "video_count": 0,
                "total_like": 0,
                "total_comment": 0,
                "total_share": 0,
                "total_collect": 0,
                "tags_set": set(),
            }
        a = author_stats[aid]
        a["video_count"] += 1
        a["total_like"] += row["点赞数"]
        a["total_comment"] += row["评论数"]
        a["total_share"] += row["分享数"]
        a["total_collect"] += row["收藏数"]
        tags = re.findall(r'#(\S+)', row["话题标签"])
        a["tags_set"].update(tags)

    for a in author_stats.values():
        a["tags_str"] = " ".join(f"#{t}" for t in list(a["tags_set"])[:10])
        a["avg_like"] = a["total_like"] // max(a["video_count"], 1)
        a["avg_comment"] = a["total_comment"] // max(a["video_count"], 1)

    author_rows = sorted(author_stats.values(), key=lambda x: x["follower_count"], reverse=True)

    # ── 非洲相关性分析 ──
    africa_kw = ['非洲','出海','跨境','海外','外贸','非洲人','非洲国家','撒哈拉',
                 '肯尼亚','尼日利亚','埃塞','加纳','坦桑','莫桑','乌干达',
                 '刚果','安哥拉','卢旺达','赞比亚','津巴布韦','喀麦隆','塞内加尔',
                 '苏丹','阿尔及利亚','埃及','摩洛哥','利比亚','突尼斯',
                 '东非','西非','北非','中非','Afrika','Africa']

    # 按作者汇总所有信号（昵称、话题、描述）
    author_signals = {}
    for row in rows:
        aid = row["作者ID"]
        if aid not in author_signals:
            author_signals[aid] = {
                "nickname": row["作者昵称"],
                "tags": set(),
                "desc_text": "",
            }
        author_signals[aid]["tags"].update(re.findall(r'#(\S+)', row["话题标签"]))
        author_signals[aid]["desc_text"] += " " + (row["视频描述"] or "")

    def is_africa_related(aid):
        s = author_signals.get(aid)
        if not s: return False
        # 检查昵称
        for kw in africa_kw:
            if kw in s["nickname"]: return True
        # 检查话题
        for tag in s.get("tags", set()):
            for kw in africa_kw:
                if kw in tag: return True
        # 检查描述（前2000字符）
        desc = s.get("desc_text", "")[:2000]
        for kw in africa_kw:
            if kw in desc: return True
        return False

    # 给每行标注
    for row in rows:
        row["非洲相关"] = "✓" if is_africa_related(row["作者ID"]) else "?"    # ── 写入 Excel ──
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        log("安装 openpyxl...")
        import subprocess
        subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()

    # ── Sheet1: 视频明细 ──
    ws1 = wb.active
    ws1.title = "视频明细"

    headers = ["视频链接","非洲相关","作者昵称","作者主页","作者签名","粉丝数(文本)","点赞数","评论数","分享数","收藏数","话题标签","视频描述","时长","播放量(文本)","来源关键词","热评1","热评1点赞","热评2","热评2点赞"]
    ws1.append(headers)

    # 标题行样式
    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    hdr_font = Font(bold=True, color="FFFFFF", size=11)
    for col in range(1, len(headers)+1):
        cell = ws1.cell(row=1, column=col)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in rows:
        ws1.append([
            row["视频链接"],
            row["非洲相关"],
            row["作者昵称"],
            row["作者主页"],
            row["作者签名"],
            row["粉丝数(文本)"],
            row["点赞数"],
            row["评论数"],
            row["分享数"],
            row["收藏数"],
            row["话题标签"],
            row["视频描述"],
            row["时长"],
            row["播放量(文本)"],
            row["来源关键词"],
            row["热评1"],
            row["热评1点赞"],
            row["热评2"],
            row["热评2点赞"],
        ])

    # 数字列居中
    center_cols = [7,8,9,10,17]  # 点赞/评论/分享/收藏/热评点赞
    for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row):
        for i, cell in enumerate(row, 1):
            if i in center_cols:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            # 数字格式
            if i in [7,8,9,10,17]:
                cell.number_format = '#,##0'

    # 列宽
    col_widths = [50, 8, 18, 40, 45, 14, 12, 10, 10, 10, 35, 50, 8, 14, 16, 50, 10, 50, 10]
    for i, w in enumerate(col_widths, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w
    ws1.row_dimensions[1].height = 30
    ws1.freeze_panes = "A2"

    # ── Sheet2: 作者汇总 ──
    ws2 = wb.create_sheet("作者汇总")
    hdr2 = ["作者昵称","作者主页","作者签名","作者ID","粉丝数","粉丝数(文本)","视频数","总点赞","平均点赞","总评论","总分享","总收藏","话题列表"]
    ws2.append(hdr2)
    for col in range(1, len(hdr2)+1):
        cell = ws2.cell(row=1, column=col)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for a in author_rows:
        follower = a["follower_count"]
        fc_text = format_count(follower)
        homepage = f"https://www.douyin.com/user/{a['sec_uid']}"
        ws2.append([
            a["nickname"], homepage, a["signature"], a["sec_uid"], follower, fc_text,
            a["video_count"], a["total_like"], a["avg_like"],
            a["total_comment"], a["total_share"], a["total_collect"], a["tags_str"]
        ])

    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row):
        for i, cell in enumerate(row, 1):
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if i in [5,6,7,8,9,10,11,12]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.number_format = '#,##0'

    col_widths2 = [20, 40, 45, 35, 14, 14, 10, 14, 14, 12, 12, 12, 40]
    for i, w in enumerate(col_widths2, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.row_dimensions[1].height = 30
    ws2.freeze_panes = "A2"

    # ── Sheet3: 关键词统计 ──
    ws3 = wb.create_sheet("关键词统计")
    kw_stats = Counter(row["来源关键词"] for row in rows)
    ws3.append(["来源关键词", "视频数", "总点赞", "平均点赞", "总评论"])
    for col in range(1, 6):
        cell = ws3.cell(row=1, column=col)
        cell.fill = hdr_fill; cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for kw, cnt in kw_stats.most_common():
        kw_rows = [r for r in rows if r["来源关键词"] == kw]
        total_like = sum(r["点赞数"] for r in kw_rows)
        avg_like = total_like // max(cnt, 1)
        total_cmt = sum(r["评论数"] for r in kw_rows)
        ws3.append([kw, cnt, total_like, avg_like, total_cmt])

    for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row):
        for i, cell in enumerate(row, 1):
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if i in [2,3,4,5]: cell.number_format = '#,##0'
    ws3.column_dimensions["A"].width = 20
    for i in range(2,6): ws3.column_dimensions[get_column_letter(i)].width = 14
    ws3.row_dimensions[1].height = 30
    ws3.freeze_panes = "A2"

    # ── Sheet4: 主题无关账号 ──
    ws4 = wb.create_sheet("主题无关账号")
    hdr4 = ["视频链接","作者昵称","作者主页","作者签名","粉丝数(文本)","点赞数","评论数","分享数","收藏数","来源关键词","视频描述"]
    ws4.append(hdr4)
    for col in range(1, len(hdr4)+1):
        cell = ws4.cell(row=1, column=col)
        cell.fill = PatternFill("solid", fgColor="8B0000")
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    unrelated_rows = [r for r in rows if r["非洲相关"] == "?"]
    for row in unrelated_rows:
        ws4.append([
            row["视频链接"], row["作者昵称"], row["作者主页"], row["作者签名"], row["粉丝数(文本)"],
            row["点赞数"], row["评论数"], row["分享数"], row["收藏数"],
            row["来源关键词"], row["视频描述"],
        ])

    for row in ws4.iter_rows(min_row=2, max_row=ws4.max_row):
        for i, cell in enumerate(row, 1):
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if i in [6,7,8,9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.number_format = '#,##0'

    col_w4 = [50, 18, 40, 45, 14, 12, 10, 10, 10, 16, 50]
    for i, w in enumerate(col_w4, 1):
        ws4.column_dimensions[get_column_letter(i)].width = w
    ws4.row_dimensions[1].height = 30
    ws4.freeze_panes = "A2"

    log(f"   主题无关: {len(unrelated_rows)} 条视频")

    # 保存
    out_path = WORK_DIR / "抖音非洲生意账号分析报告_v4.xlsx"
    wb.save(str(out_path))
    log(f"\n✅ Excel 已保存: {out_path}")
    log(f"   文件大小: {out_path.stat().st_size//1024} KB")
    log(f"   视频明细: {len(rows)} 条")
    log(f"   作者汇总: {len(author_rows)} 个账号")
    log(f"   关键词数: {len(kw_stats)} 个")

    # 打印摘要
    print_summary(rows, author_rows, kw_stats)

def print_summary(rows, author_rows, kw_stats):
    log("\n" + "="*60)
    log("报告摘要")
    log("="*60)
    log(f"总视频数: {len(rows)}")
    log(f"去重作者数: {len(author_rows)}")

    total_like = sum(r["点赞数"] for r in rows)
    total_cmt = sum(r["评论数"] for r in rows)
    log(f"总点赞: {total_like:,}  总评论: {total_cmt:,}")

    log("\n--- TOP 10 账号（按粉丝数）---")
    for i, a in enumerate(author_rows[:10], 1):
        fc_text = format_count(a["follower_count"])
        log(f"  {i}. {a['nickname']}  粉丝: {fc_text}  视频: {a['video_count']}  均赞: {a['avg_like']:,}")

    log("\n--- 关键词分布 ---")
    for kw, cnt in kw_stats.most_common():
        log(f"  {kw}: {cnt} 条")

    log("\n--- 互动最高的视频 TOP 5 ---")
    top_videos = sorted(rows, key=lambda r: r["点赞数"], reverse=True)[:5]
    for i, v in enumerate(top_videos, 1):
        log(f"  {i}. [{v['点赞数']:,}赞] {v['作者昵称']} - {v['视频描述'][:40]}")

if __name__ == "__main__":
    main()
